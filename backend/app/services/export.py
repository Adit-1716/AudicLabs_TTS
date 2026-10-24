import io
from datetime import timedelta

from docx import Document


def as_txt(segments: list, speakers: dict, raw: str) -> bytes:
    if segments:
        lines = []
        for s in segments:
            speaker = s.get("speaker") or speakers.get(str(s.get("speaker_id")), "Speaker")
            lines.append(f"[{_ts(s.get('start', 0))}] {speaker}: {s.get('text','')}")
        return "\n".join(lines).encode("utf-8")
    return (raw or "").encode("utf-8")


def as_srt(segments: list, raw: str) -> bytes:
    if not segments:
        segments = [{"start": 0, "end": 5, "text": raw}]
    blocks = []
    for i, s in enumerate(segments, 1):
        start = _srt(s.get("start") or 0)
        end = _srt(s.get("end") or (s.get("start") or 0) + 2)
        speaker = s.get("speaker")
        text = s.get("text") or ""
        if speaker:
            text = f"{speaker}: {text}"
        blocks.append(f"{i}\n{start} --> {end}\n{text}\n")
    return "\n".join(blocks).encode("utf-8")


def as_docx(title: str, segments: list, raw: str) -> bytes:
    doc = Document()
    doc.add_heading(title, 1)
    if segments:
        for s in segments:
            p = doc.add_paragraph()
            run = p.add_run(f"{s.get('speaker','Speaker')}  {_ts(s.get('start', 0))}")
            run.bold = True
            doc.add_paragraph(s.get("text") or "")
    else:
        doc.add_paragraph(raw or "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _ts(seconds: float) -> str:
    seconds = float(seconds or 0)
    m, s = divmod(int(seconds), 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def _srt(seconds: float) -> str:
    td = timedelta(seconds=float(seconds or 0))
    total = int(td.total_seconds())
    ms = int((float(seconds or 0) - total) * 1000)
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"
